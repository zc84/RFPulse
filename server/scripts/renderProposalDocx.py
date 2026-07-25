import json
import math
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROBOTO = 'Roboto'
MONO = 'Consolas'
CONTENT_MARKER = '{{PROPOSAL_CONTENT}}'
PREFERRED_TABLE_STYLE = 'Plain Table 3'

# Andersen brand palette (andersenlab.com)
ACCENT = 'FFDB00'        # signature yellow
INK = '020D1C'           # near-black for headings and body
SLATE = '556170'         # secondary text
SLATE_DARK = '25303E'    # table header fill
PAPER = 'F5F6F7'         # zebra / code background
BORDER = 'D9DCE1'        # subtle rules
LINK = '004D74'          # hyperlink blue from the brand's dark palette

HEADING_SIZES = {'Title': 26, 'Heading 1': 16, 'Heading 2': 13, 'Heading 3': 11.5}
HEADING_COLORS = {'Title': INK, 'Heading 1': INK, 'Heading 2': INK, 'Heading 3': SLATE_DARK}

TECHNICAL_SECTION_PATTERNS = [
  re.compile(r'^\s*#{1,6}\s+.*\b(proposed architecture|solution architecture|technical solution|proposed solution|solution overview|architecture)\b.*$', re.IGNORECASE),
]
TIMELINE_SECTION_PATTERN = re.compile(r'^\s*#{1,6}\s+.*\b(timeline|implementation plan|delivery plan|roadmap|project schedule|schedule)\b.*$', re.IGNORECASE)
HEADING_PATTERN = re.compile(r'^(#{1,6})\s+(.*)$')
PHASE_LINE_PATTERN = re.compile(r'\bphase\s+([ivx]+|\d+)\b', re.IGNORECASE)
DURATION_PATTERN = re.compile(r'(\d+(?:\.\d+)?)\s*(day|days|week|weeks|month|months)\b', re.IGNORECASE)
HR_PATTERN = re.compile(r'^\s*(-{3,}|\*{3,}|_{3,})\s*$')
GANTT_SENTINEL = '@@RFPULSE-GANTT@@'
DIAGRAM_SENTINEL = '@@RFPULSE-DIAGRAM@@'

INLINE_TOKEN = re.compile(
  r'(?P<code>`[^`]+`)'
  r'|(?P<bold>(\*\*|__)(?:(?!\3).)+\3)'
  r'|(?P<italic>(\*|_)(?:(?!\5).)+\5)'
  r'|(?P<image>!\[[^\]]*\]\([^)]+\))'
  r'|(?P<link>\[[^\]]+\]\([^)]+\))'
)


def hex_color(value):
  return RGBColor.from_string(value)


def set_run_font(run, size=10.5, bold=None, italic=None, color=INK, name=ROBOTO):
  run.font.name = name
  rPr = run._element.get_or_add_rPr()
  rFonts = rPr.rFonts
  if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.append(rFonts)
  rFonts.set(qn('w:eastAsia'), name)
  run.font.size = Pt(size)
  if color:
    run.font.color.rgb = hex_color(color)
  if bold is not None:
    run.bold = bold
  if italic is not None:
    run.italic = italic


def set_style_font(style, size, bold=False, color=INK):
  style.font.name = ROBOTO
  rPr = style._element.get_or_add_rPr()
  rFonts = rPr.rFonts
  if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.append(rFonts)
  rFonts.set(qn('w:eastAsia'), ROBOTO)
  style.font.size = Pt(size)
  style.font.bold = bold
  if color:
    style.font.color.rgb = hex_color(color)


def set_paragraph_shading(paragraph, fill):
  pPr = paragraph._p.get_or_add_pPr()
  shd = OxmlElement('w:shd')
  shd.set(qn('w:val'), 'clear')
  shd.set(qn('w:fill'), fill)
  pPr.append(shd)


def set_paragraph_bottom_border(paragraph, color, size=8, space=4):
  pPr = paragraph._p.get_or_add_pPr()
  pBdr = OxmlElement('w:pBdr')
  bottom = OxmlElement('w:bottom')
  bottom.set(qn('w:val'), 'single')
  bottom.set(qn('w:sz'), str(size))
  bottom.set(qn('w:space'), str(space))
  bottom.set(qn('w:color'), color)
  pBdr.append(bottom)
  pPr.append(pBdr)


def set_cell_shading(cell, fill):
  tcPr = cell._tc.get_or_add_tcPr()
  shd = OxmlElement('w:shd')
  shd.set(qn('w:val'), 'clear')
  shd.set(qn('w:fill'), fill)
  tcPr.append(shd)


def set_repeat_table_header(row):
  trPr = row._tr.get_or_add_trPr()
  tblHeader = OxmlElement('w:tblHeader')
  tblHeader.set(qn('w:val'), 'true')
  trPr.append(tblHeader)


def add_field(paragraph, instruction, size=9, color=SLATE):
  begin = OxmlElement('w:r')
  fld_begin = OxmlElement('w:fldChar')
  fld_begin.set(qn('w:fldCharType'), 'begin')
  begin.append(fld_begin)
  paragraph._p.append(begin)

  instr_run = OxmlElement('w:r')
  instr = OxmlElement('w:instrText')
  instr.set(qn('xml:space'), 'preserve')
  instr.text = instruction
  instr_run.append(instr)
  paragraph._p.append(instr_run)

  end = OxmlElement('w:r')
  fld_end = OxmlElement('w:fldChar')
  fld_end.set(qn('w:fldCharType'), 'end')
  end.append(fld_end)
  paragraph._p.append(end)

  # Give the surrounding literal runs the same look as the field result.
  for run in paragraph.runs:
    if run.font.size is None:
      set_run_font(run, size=size, color=color)


def enable_update_fields_on_open(doc):
  try:
    settings = doc.settings.element
  except (AttributeError, KeyError):
    return
  if settings.find(qn('w:updateFields')) is None:
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)


def add_hyperlink(paragraph, text, url, size=10.5):
  try:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
  except Exception:
    run = paragraph.add_run(f'{text} ({url})')
    set_run_font(run, size=size)
    return
  hyperlink = OxmlElement('w:hyperlink')
  hyperlink.set(qn('r:id'), r_id)
  run_el = OxmlElement('w:r')
  rPr = OxmlElement('w:rPr')
  rFonts = OxmlElement('w:rFonts')
  rFonts.set(qn('w:ascii'), ROBOTO)
  rFonts.set(qn('w:hAnsi'), ROBOTO)
  rPr.append(rFonts)
  sz = OxmlElement('w:sz')
  sz.set(qn('w:val'), str(int(size * 2)))
  rPr.append(sz)
  color = OxmlElement('w:color')
  color.set(qn('w:val'), LINK)
  rPr.append(color)
  underline = OxmlElement('w:u')
  underline.set(qn('w:val'), 'single')
  rPr.append(underline)
  run_el.append(rPr)
  text_el = OxmlElement('w:t')
  text_el.set(qn('xml:space'), 'preserve')
  text_el.text = text
  run_el.append(text_el)
  hyperlink.append(run_el)
  paragraph._p.append(hyperlink)


def clean_inline(text):
  """Flatten inline markdown to plain text (used where rich runs are impossible)."""
  text = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)', r'\1', text)
  text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
  text = re.sub(r'(\*\*|__)(.*?)\1', r'\2', text)
  text = re.sub(r'(\*|_)(.*?)\1', r'\2', text)
  text = re.sub(r'`([^`]*)`', r'\1', text)
  return text.strip()


def add_inline_runs(paragraph, text, size=10.5, bold=False, italic=False, color=INK):
  """Render markdown inline formatting (**bold**, *italic*, `code`, [links](url)) as styled runs."""
  pos = 0
  for match in INLINE_TOKEN.finditer(text):
    if match.start() > pos:
      run = paragraph.add_run(text[pos:match.start()])
      set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    token = match.group(0)
    if match.group('code'):
      run = paragraph.add_run(token[1:-1])
      set_run_font(run, size=max(size - 1, 8), bold=bold, italic=italic, color=SLATE_DARK, name=MONO)
    elif match.group('bold'):
      inner = token[2:-2]
      # Support one level of nested italic inside bold.
      add_inline_runs(paragraph, inner, size=size, bold=True, italic=italic, color=color)
    elif match.group('italic'):
      inner = token[1:-1]
      run = paragraph.add_run(clean_inline(inner))
      set_run_font(run, size=size, bold=bold, italic=True, color=color)
    elif match.group('image'):
      alt = re.match(r'!\[([^\]]*)\]', token).group(1)
      if alt:
        run = paragraph.add_run(alt)
        set_run_font(run, size=size, bold=bold, italic=True, color=SLATE)
    elif match.group('link'):
      link_match = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
      label, url = link_match.group(1), link_match.group(2)
      if re.match(r'^https?://', url):
        add_hyperlink(paragraph, clean_inline(label), url, size=size)
      else:
        run = paragraph.add_run(clean_inline(label))
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    pos = match.end()
  if pos < len(text):
    run = paragraph.add_run(text[pos:])
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)


def configure_document_layout(doc):
  for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

  if 'Normal' in doc.styles:
    normal = doc.styles['Normal'].paragraph_format
    normal.space_before = Pt(0)
    normal.space_after = Pt(8)
    normal.line_spacing = 1.15

  for style_name, before, after, keep_with_next in [
    ('Title', 0, 14, False),
    ('Heading 1', 22, 10, True),
    ('Heading 2', 16, 8, True),
    ('Heading 3', 12, 6, True),
  ]:
    if style_name in doc.styles:
      pf = doc.styles[style_name].paragraph_format
      pf.space_before = Pt(before)
      pf.space_after = Pt(after)
      pf.keep_with_next = keep_with_next

  for style_name in ['List Bullet', 'List Number']:
    if style_name in doc.styles:
      pf = doc.styles[style_name].paragraph_format
      pf.left_indent = Inches(0.25)
      pf.space_before = Pt(0)
      pf.space_after = Pt(4)
      pf.line_spacing = 1.1


def configure_styles(doc):
  for style_name, size, bold in [
    ('Normal', 10.5, False),
    ('Title', HEADING_SIZES['Title'], True),
    ('Heading 1', HEADING_SIZES['Heading 1'], True),
    ('Heading 2', HEADING_SIZES['Heading 2'], True),
    ('Heading 3', HEADING_SIZES['Heading 3'], True),
    ('List Bullet', 10.5, False),
    ('List Number', 10.5, False),
  ]:
    if style_name in doc.styles:
      set_style_font(doc.styles[style_name], size, bold, color=HEADING_COLORS.get(style_name, INK))
  configure_document_layout(doc)


def has_style(doc, style_name):
  try:
    doc.styles[style_name]
    return True
  except KeyError:
    return False


def add_paragraph_text(doc, text, style='Normal', align=None, bold=False, italic=False, size=None, color=INK):
  if not has_style(doc, style):
    style = 'Normal'
  para = doc.add_paragraph(style=style)
  if align is not None:
    para.alignment = align
  font_size = size if size is not None else HEADING_SIZES.get(style, 10.5)
  add_inline_runs(para, text, size=font_size, bold=bold or style in HEADING_SIZES, italic=italic,
                  color=HEADING_COLORS.get(style, color))
  if style in ('List Bullet', 'List Number'):
    para.paragraph_format.left_indent = Inches(0.25)
  return para


def is_table_separator(line):
  return bool(re.match(r'^\s*\|?[\s:-]+(\|[\s:-]+)+\|?\s*$', line))


def parse_table_row(line):
  return [cell.replace(r'\|', '|').strip() for cell in line.strip().strip('|').split('|')]


def preferred_table_style_name(doc):
  if has_style(doc, PREFERRED_TABLE_STYLE):
    return PREFERRED_TABLE_STYLE
  if has_style(doc, 'Table Grid'):
    return 'Table Grid'
  return None


def apply_table_style(table, style_name):
  if not style_name:
    return
  table.style = style_name


def style_data_table(table, zebra=True):
  """Andersen-branded table look: dark header row with white text, zebra body rows."""
  if not table.rows:
    return
  header = table.rows[0]
  set_repeat_table_header(header)
  for cell in header.cells:
    set_cell_shading(cell, SLATE_DARK)
    for para in cell.paragraphs:
      for run in para.runs:
        set_run_font(run, size=9.5, bold=True, color='FFFFFF')
  if zebra:
    for r_index, row in enumerate(table.rows[1:], start=1):
      if r_index % 2 == 0:
        for cell in row.cells:
          set_cell_shading(cell, PAPER)


def normalize_table_styles(doc):
  style_name = preferred_table_style_name(doc)
  if not style_name:
    return
  for table in doc.tables:
    apply_table_style(table, style_name)


def clear_paragraph_content(paragraph):
  p = paragraph._p
  for child in list(p):
    p.remove(child)


# --- Numbered list restart support -------------------------------------------------

DECIMAL_ABSTRACT_NUM_ID = 991


def get_numbering_element(doc):
  try:
    return doc.part.numbering_part.element
  except (NotImplementedError, KeyError, AttributeError):
    return None


def ensure_decimal_abstract_num(numbering):
  for el in numbering.findall(qn('w:abstractNum')):
    if el.get(qn('w:abstractNumId')) == str(DECIMAL_ABSTRACT_NUM_ID):
      return
  abstract = OxmlElement('w:abstractNum')
  abstract.set(qn('w:abstractNumId'), str(DECIMAL_ABSTRACT_NUM_ID))
  for level in range(3):
    lvl = OxmlElement('w:lvl')
    lvl.set(qn('w:ilvl'), str(level))
    start = OxmlElement('w:start')
    start.set(qn('w:val'), '1')
    lvl.append(start)
    fmt = OxmlElement('w:numFmt')
    fmt.set(qn('w:val'), 'decimal')
    lvl.append(fmt)
    text = OxmlElement('w:lvlText')
    text.set(qn('w:val'), f'%{level + 1}.')
    lvl.append(text)
    jc = OxmlElement('w:lvlJc')
    jc.set(qn('w:val'), 'left')
    lvl.append(jc)
    pPr = OxmlElement('w:pPr')
    ind = OxmlElement('w:ind')
    ind.set(qn('w:left'), str(432 * (level + 1)))
    ind.set(qn('w:hanging'), '216')
    pPr.append(ind)
    lvl.append(pPr)
    abstract.append(lvl)
  nums = numbering.findall(qn('w:num'))
  if nums:
    nums[0].addprevious(abstract)
  else:
    numbering.append(abstract)


def create_numbering_instance(doc):
  """Create a fresh numbering instance so each numbered list restarts at 1."""
  numbering = get_numbering_element(doc)
  if numbering is None:
    return None
  ensure_decimal_abstract_num(numbering)
  existing = [int(num.get(qn('w:numId'))) for num in numbering.findall(qn('w:num'))]
  next_id = max(existing, default=0) + 1
  num = OxmlElement('w:num')
  num.set(qn('w:numId'), str(next_id))
  ref = OxmlElement('w:abstractNumId')
  ref.set(qn('w:val'), str(DECIMAL_ABSTRACT_NUM_ID))
  num.append(ref)
  numbering.append(num)
  return next_id


def apply_paragraph_numbering(paragraph, num_id, level=0):
  pPr = paragraph._p.get_or_add_pPr()
  numPr = OxmlElement('w:numPr')
  ilvl = OxmlElement('w:ilvl')
  ilvl.set(qn('w:val'), str(level))
  numPr.append(ilvl)
  numId = OxmlElement('w:numId')
  numId.set(qn('w:val'), str(num_id))
  numPr.append(numId)
  pPr.append(numPr)


def bullet_indent_level(raw_line):
  indent = len(raw_line) - len(raw_line.lstrip(' '))
  if indent >= 4:
    return 2
  if indent >= 2:
    return 1
  return 0


BULLET_STYLES = {0: 'List Bullet', 1: 'List Bullet 2', 2: 'List Bullet 3'}


def add_code_block(doc, code_lines):
  for line in code_lines or ['']:
    para = doc.add_paragraph()
    set_paragraph_shading(para, PAPER)
    para.paragraph_format.left_indent = Inches(0.2)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(line if line else ' ')
    set_run_font(run, size=9, color=SLATE_DARK, name=MONO)
  doc.add_paragraph('')


def add_horizontal_rule(doc):
  para = doc.add_paragraph()
  para.paragraph_format.space_before = Pt(6)
  para.paragraph_format.space_after = Pt(6)
  set_paragraph_bottom_border(para, BORDER, size=6, space=1)


def add_markdown(doc, markdown, inline_diagram_images=None):
  lines = markdown.splitlines()
  paragraph_buffer = []
  table_style_name = preferred_table_style_name(doc)
  active_num_id = None
  diagram_lookup = {
    item.get('placeholder'): item
    for item in (inline_diagram_images or [])
    if item.get('placeholder')
  }

  def flush_paragraph():
    nonlocal paragraph_buffer
    if paragraph_buffer:
      text = ' '.join(part.strip() for part in paragraph_buffer if part.strip())
      if text:
        add_paragraph_text(doc, text)
      paragraph_buffer = []

  i = 0
  while i < len(lines):
    line = lines[i].rstrip('\n')
    stripped = line.strip()

    if not stripped:
      flush_paragraph()
      active_num_id = None
      i += 1
      continue

    if stripped.startswith(GANTT_SENTINEL):
      flush_paragraph()
      try:
        phases = json.loads(stripped[len(GANTT_SENTINEL):])
      except ValueError:
        phases = []
      if phases:
        add_gantt_table(doc, phases)
      i += 1
      continue

    if stripped.startswith(DIAGRAM_SENTINEL):
      flush_paragraph()
      active_num_id = None
      diagram = diagram_lookup.get(stripped)
      if diagram:
        add_single_diagram(doc, diagram)
      i += 1
      continue

    if HR_PATTERN.match(stripped):
      flush_paragraph()
      active_num_id = None
      add_horizontal_rule(doc)
      i += 1
      continue

    heading = HEADING_PATTERN.match(stripped)
    if heading:
      flush_paragraph()
      active_num_id = None
      level = min(len(heading.group(1)), 3)
      style = f'Heading {level}'
      para = add_paragraph_text(doc, heading.group(2), style=style)
      para.paragraph_format.space_before = Pt(22 if level == 1 else 16 if level == 2 else 12)
      para.paragraph_format.space_after = Pt(10 if level == 1 else 8 if level == 2 else 6)
      para.paragraph_format.keep_with_next = True
      if level == 1:
        set_paragraph_bottom_border(para, ACCENT, size=12, space=3)
      i += 1
      continue

    if stripped.startswith('```'):
      flush_paragraph()
      active_num_id = None
      code_lines = []
      i += 1
      while i < len(lines) and not lines[i].strip().startswith('```'):
        code_lines.append(lines[i].rstrip('\n'))
        i += 1
      add_code_block(doc, code_lines)
      i += 1
      continue

    if stripped.startswith('|') and '|' in stripped[1:]:
      flush_paragraph()
      active_num_id = None
      table_rows = []
      while i < len(lines):
        row_line = lines[i].rstrip('\n')
        row_stripped = row_line.strip()
        if not row_stripped.startswith('|') or is_table_separator(row_stripped):
          if is_table_separator(row_stripped):
            i += 1
            continue
          break
        table_rows.append(parse_table_row(row_stripped))
        i += 1
      if table_rows:
        max_cols = max(len(row) for row in table_rows)
        table = doc.add_table(rows=len(table_rows), cols=max_cols)
        apply_table_style(table, table_style_name)
        for r_index, row in enumerate(table_rows):
          for c_index in range(max_cols):
            text = row[c_index] if c_index < len(row) else ''
            cell_para = table.rows[r_index].cells[c_index].paragraphs[0]
            cell_para.text = ''
            add_inline_runs(cell_para, text, size=9.5, bold=(r_index == 0),
                            color='FFFFFF' if r_index == 0 else INK)
        style_data_table(table)
        doc.add_paragraph('')
      continue

    bullet = re.match(r'^\s*[-*]\s+(.*)$', line)
    if bullet:
      flush_paragraph()
      active_num_id = None
      level = bullet_indent_level(line)
      style = BULLET_STYLES.get(level, 'List Bullet')
      para = add_paragraph_text(doc, bullet.group(1), style=style if has_style(doc, style) else 'List Bullet')
      para.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
      i += 1
      continue

    numbered = re.match(r'^\s*\d+\.\s+(.*)$', stripped)
    if numbered:
      flush_paragraph()
      if active_num_id is None:
        active_num_id = create_numbering_instance(doc)
      para = add_paragraph_text(doc, numbered.group(1), style='List Number')
      if active_num_id is not None:
        apply_paragraph_numbering(para, active_num_id)
      i += 1
      continue

    quote = re.match(r'^\s*>\s+(.*)$', stripped)
    if quote:
      flush_paragraph()
      active_num_id = None
      para = doc.add_paragraph()
      para.paragraph_format.left_indent = Inches(0.3)
      set_paragraph_shading(para, PAPER)
      add_inline_runs(para, quote.group(1), size=10.5, italic=True, color=SLATE)
      i += 1
      continue

    paragraph_buffer.append(stripped)
    i += 1

  flush_paragraph()


def add_diagrams(doc, diagrams):
  if not diagrams:
    return
  for index, diagram in enumerate(diagrams, start=1):
    image_path = diagram['path']
    if not Path(image_path).exists():
      raise FileNotFoundError(f'Diagram image not found: {image_path}')
    title = diagram.get('title') or 'Diagram'
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(image_path, width=Inches(6.5))
    if title:
      caption = doc.add_paragraph()
      caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
      cap_run = caption.add_run(f'Figure {index}. {clean_inline(title)}')
      set_run_font(cap_run, size=9.5, italic=True, color=SLATE)
    description = clean_inline(diagram.get('description') or '')
    if description:
      summary = doc.add_paragraph()
      summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
      sum_run = summary.add_run(description)
      set_run_font(sum_run, size=9, color=SLATE)
    doc.add_paragraph('')


def add_single_diagram(doc, diagram, index_hint=None):
  if not diagram:
    return
  image_path = diagram['path']
  if not Path(image_path).exists():
    raise FileNotFoundError(f'Diagram image not found: {image_path}')
  title = diagram.get('title') or 'Diagram'
  para = doc.add_paragraph()
  para.alignment = WD_ALIGN_PARAGRAPH.CENTER
  run = para.add_run()
  run.add_picture(image_path, width=Inches(6.5))
  if title:
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = f'Figure {index_hint}. ' if index_hint else ''
    cap_run = caption.add_run(f'{prefix}{clean_inline(title)}')
    set_run_font(cap_run, size=9.5, italic=True, color=SLATE)
  description = clean_inline(diagram.get('description') or '')
  if description:
    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sum_run = summary.add_run(description)
    set_run_font(sum_run, size=9, color=SLATE)
  doc.add_paragraph('')


def trim_architecture_narrative_heading(markdown):
  lines = markdown.splitlines()
  if not lines:
    return markdown

  first_nonempty = None
  for idx, line in enumerate(lines):
    if line.strip():
      first_nonempty = idx
      break

  if first_nonempty is None:
    return markdown

  heading = HEADING_PATTERN.match(lines[first_nonempty].strip())
  if not heading:
    return markdown

  heading_text = heading.group(2).strip().lower()
  if heading_text not in ('architecture narrative', 'diagram narrative', 'architecture diagrams narrative'):
    return markdown

  start = first_nonempty + 1
  while start < len(lines) and not lines[start].strip():
    start += 1
  return '\n'.join(lines[start:]).strip()


def duration_to_weeks(duration_matches):
  total_weeks = 0.0
  for value, unit in duration_matches:
    amount = float(value)
    normalized = unit.lower()
    if normalized.startswith('day'):
      total_weeks += amount / 5.0
    elif normalized.startswith('week'):
      total_weeks += amount
    elif normalized.startswith('month'):
      total_weeks += amount * 4.0
  return total_weeks


def extract_timeline_phases(section_lines):
  phases = []
  seen = set()
  for raw in section_lines:
    line = raw.strip()
    if not line or line.startswith('|'):
      continue
    phase_match = PHASE_LINE_PATTERN.search(line)
    durations = DURATION_PATTERN.findall(line)
    if not phase_match or not durations:
      continue
    phase_label = phase_match.group(0).title()
    cleaned = clean_inline(re.sub(r'^[-*]\s+', '', line))
    name = cleaned.split(':', 1)[0] if ':' in cleaned else phase_label
    weeks = duration_to_weeks(durations)
    if weeks <= 0:
      continue
    key = name.lower()
    if key in seen:
      continue
    seen.add(key)
    phases.append({
      'name': name,
      'durationLabel': ', '.join([f"{value} {unit}" for value, unit in durations]),
      'weeks': weeks,
    })
  return phases


def add_gantt_table(doc, phases):
  """Render a shaded-cell delivery timeline instead of a text-art bar chart."""
  total_weeks = sum(max(1, int(math.ceil(phase['weeks']))) for phase in phases)
  col_count = total_weeks

  title = add_paragraph_text(doc, 'Delivery Timeline', style='Heading 3')
  title.paragraph_format.space_before = Pt(12)

  table = doc.add_table(rows=len(phases) + 2, cols=2 + col_count)
  apply_table_style(table, preferred_table_style_name(doc))
  month_header = table.rows[0]
  week_header = table.rows[1]

  for first_col, label in ((0, 'Phase'), (1, 'Duration')):
    merged = month_header.cells[first_col].merge(week_header.cells[first_col])
    para = merged.paragraphs[0]
    para.text = ''
    run = para.add_run(label)
    set_run_font(run, size=9, bold=True, color='FFFFFF')
    set_cell_shading(merged, SLATE_DARK)

  week_index = 0
  month_number = 1
  while week_index < col_count:
    start_col = 2 + week_index
    end_col = 2 + min(week_index + 3, col_count - 1)
    cell = month_header.cells[start_col]
    if end_col > start_col:
      cell = cell.merge(month_header.cells[end_col])
    para = cell.paragraphs[0]
    para.text = ''
    run = para.add_run(f'M{month_number}')
    set_run_font(run, size=8.5, bold=True, color='FFFFFF')
    set_cell_shading(cell, SLATE_DARK)
    month_number += 1
    week_index += 4

  for week in range(col_count):
    cell = week_header.cells[2 + week]
    para = cell.paragraphs[0]
    para.text = ''
    run = para.add_run(f'W{(week % 4) + 1}')
    set_run_font(run, size=7.5, bold=True, color='FFFFFF')
    set_cell_shading(cell, SLATE_DARK)

  set_repeat_table_header(month_header)
  set_repeat_table_header(week_header)

  start_week = 0  # zero-based
  for p_index, phase in enumerate(phases, start=2):
    duration = max(1, int(math.ceil(phase['weeks'])))
    row = table.rows[p_index]
    name_para = row.cells[0].paragraphs[0]
    name_para.text = ''
    name_run = name_para.add_run(clean_inline(phase['name']))
    set_run_font(name_run, size=9, bold=True)
    duration_para = row.cells[1].paragraphs[0]
    duration_para.text = ''
    duration_run = duration_para.add_run(phase['durationLabel'])
    set_run_font(duration_run, size=8.5, color=SLATE)
    for week in range(start_week, min(start_week + duration, col_count)):
      set_cell_shading(row.cells[2 + week], ACCENT)
    start_week += duration

  widths = [Inches(1.7), Inches(1.1)] + [Inches(max(4.8 / max(col_count, 1), 0.12))] * col_count
  for row in table.rows:
    for c_index, cell in enumerate(row.cells):
      if c_index < len(widths):
        cell.width = widths[c_index]
  doc.add_paragraph('')


def inject_timeline_gantt(markdown):
  lines = markdown.splitlines()
  if any('gantt view' in line.lower() or GANTT_SENTINEL in line for line in lines):
    return markdown

  insertions = []
  i = 0
  while i < len(lines):
    heading = HEADING_PATTERN.match(lines[i].strip())
    if not heading:
      i += 1
      continue

    if not TIMELINE_SECTION_PATTERN.match(lines[i].strip()):
      i += 1
      continue

    current_level = len(heading.group(1))
    j = i + 1
    while j < len(lines):
      next_heading = HEADING_PATTERN.match(lines[j].strip())
      if next_heading and len(next_heading.group(1)) <= current_level:
        break
      j += 1

    section_lines = lines[i + 1:j]
    phases = extract_timeline_phases(section_lines)
    if len(phases) >= 2:
      insertions.append((j, ['', GANTT_SENTINEL + json.dumps(phases), '']))
    i = j

  if not insertions:
    return markdown

  for index, block in sorted(insertions, key=lambda item: item[0], reverse=True):
    lines[index:index] = block

  return '\n'.join(lines)


def split_markdown_for_diagrams(markdown):
  lines = markdown.splitlines()
  for idx, line in enumerate(lines):
    if any(pattern.match(line.strip()) for pattern in TECHNICAL_SECTION_PATTERNS):
      before = '\n'.join(lines[: idx + 1]).strip()
      after = '\n'.join(lines[idx + 1 :]).strip()
      return before, after
  return markdown.strip(), ''


def split_markdown_for_timeline(markdown):
  lines = markdown.splitlines()
  for idx, line in enumerate(lines):
    if TIMELINE_SECTION_PATTERN.match(line.strip()):
      before = '\n'.join(lines[: idx + 1]).strip()
      after = '\n'.join(lines[idx + 1 :]).strip()
      return before, after
  return markdown.strip(), ''


def count_toc_headings(markdown):
  count = 0
  for line in markdown.splitlines():
    heading = HEADING_PATTERN.match(line.strip())
    if heading and len(heading.group(1)) <= 2:
      count += 1
  return count


def add_page_break(doc):
  para = doc.add_paragraph()
  run = para.add_run()
  br = OxmlElement('w:br')
  br.set(qn('w:type'), 'page')
  run._element.append(br)


def add_cover_page(doc, title, meta):
  bar = doc.add_paragraph()
  set_paragraph_shading(bar, ACCENT)
  bar.paragraph_format.space_after = Pt(0)
  bar_run = bar.add_run(' ')
  set_run_font(bar_run, size=6)

  for _ in range(6):
    spacer = doc.add_paragraph('')
    spacer.paragraph_format.space_after = Pt(12)

  title_para = doc.add_paragraph()
  title_run = title_para.add_run(clean_inline(title))
  set_run_font(title_run, size=30, bold=True, color=INK)
  title_para.paragraph_format.space_after = Pt(6)
  set_paragraph_bottom_border(title_para, ACCENT, size=18, space=8)

  client_name = (meta or {}).get('clientName')
  if client_name:
    client_para = doc.add_paragraph()
    client_para.paragraph_format.space_before = Pt(18)
    client_run = client_para.add_run(f'Prepared for {client_name}')
    set_run_font(client_run, size=13, color=SLATE_DARK)

  date_label = (meta or {}).get('date')
  if date_label:
    date_para = doc.add_paragraph()
    date_run = date_para.add_run(date_label)
    set_run_font(date_run, size=11, color=SLATE)

  for _ in range(10):
    spacer = doc.add_paragraph('')
    spacer.paragraph_format.space_after = Pt(12)

  note = doc.add_paragraph()
  note_run = note.add_run('Confidential — prepared for evaluation purposes only.')
  set_run_font(note_run, size=9, italic=True, color=SLATE)

  add_page_break(doc)


def add_table_of_contents(doc):
  label = doc.add_paragraph()
  label_run = label.add_run('Contents')
  set_run_font(label_run, size=16, bold=True, color=INK)
  label.paragraph_format.space_after = Pt(10)
  set_paragraph_bottom_border(label, ACCENT, size=12, space=3)

  toc_para = doc.add_paragraph()
  add_field(toc_para, r'TOC \o "1-2" \h \z \u', size=10.5, color=INK)
  enable_update_fields_on_open(doc)
  add_page_break(doc)


def add_footer_with_page_numbers(doc, title):
  section = doc.sections[0]
  try:
    section.different_first_page_header_footer = True
  except AttributeError:
    pass
  footer = section.footer
  para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
  clear_paragraph_content(para)
  usable_width = section.page_width - section.left_margin - section.right_margin
  para.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
  title_run = para.add_run(clean_inline(title))
  set_run_font(title_run, size=8.5, color=SLATE)
  para.add_run('\t')
  page_label = para.add_run('Page ')
  set_run_font(page_label, size=8.5, color=SLATE)
  add_field(para, 'PAGE', size=8.5)
  of_label = para.add_run(' of ')
  set_run_font(of_label, size=8.5, color=SLATE)
  add_field(para, 'NUMPAGES', size=8.5)
  for run in para.runs:
    if run.font.size is None:
      set_run_font(run, size=8.5, color=SLATE)


def build_content(doc, title, markdown, diagrams, timeline_diagrams, inline_diagram_images, include_cover=False, include_toc=False, meta=None):
  if not timeline_diagrams:
    markdown = inject_timeline_gantt(markdown)

  if include_cover:
    add_cover_page(doc, title, meta)
  else:
    title_para = add_paragraph_text(doc, title, style='Title')
    set_paragraph_bottom_border(title_para, ACCENT, size=18, space=6)
    doc.add_paragraph('')

  if include_toc and count_toc_headings(markdown) >= 3:
    add_table_of_contents(doc)

  remaining = markdown

  if diagrams:
    before, remaining = split_markdown_for_diagrams(remaining)
    add_markdown(doc, before, inline_diagram_images)
    add_diagrams(doc, diagrams)
    remaining = trim_architecture_narrative_heading(remaining)

  if timeline_diagrams:
    before, after = split_markdown_for_timeline(remaining)
    add_markdown(doc, before, inline_diagram_images)
    add_diagrams(doc, timeline_diagrams)
    add_markdown(doc, after, inline_diagram_images)
  else:
    add_markdown(doc, remaining, inline_diagram_images)


def create_content_document(title, markdown, diagrams, timeline_diagrams, inline_diagram_images, meta=None):
  doc = Document()
  doc._body.clear_content()
  configure_styles(doc)
  build_content(doc, title, markdown, diagrams, timeline_diagrams, inline_diagram_images, include_cover=True, include_toc=True, meta=meta)
  add_footer_with_page_numbers(doc, title)
  normalize_table_styles(doc)
  return doc


def apply_template_placeholders(doc, replacements):
  def replace_in_paragraph(paragraph):
    if not any(key in paragraph.text for key in replacements):
      return
    for run in paragraph.runs:
      for key, value in replacements.items():
        if key in run.text:
          run.text = run.text.replace(key, value)
    # Placeholder split across runs: collapse to plain text as a fallback.
    remaining = [key for key in replacements if key in paragraph.text]
    if remaining:
      text = paragraph.text
      for key, value in replacements.items():
        text = text.replace(key, value)
      clear_paragraph_content(paragraph)
      run = paragraph.add_run(text)
      set_run_font(run, size=11)

  for paragraph in doc.paragraphs:
    replace_in_paragraph(paragraph)
  for table in doc.tables:
    for row in table.rows:
      for cell in row.cells:
        for paragraph in cell.paragraphs:
          replace_in_paragraph(paragraph)


def render_into_template(template_path, title, markdown, diagrams, timeline_diagrams, inline_diagram_images, meta=None):
  doc = Document(template_path)
  # Keep the template's original section layout intact (cover margins/positioning,
  # custom spacings, anchored objects). Content is built directly inside the template
  # document so style, numbering, and relationship references stay valid.

  replacements = {
    '{{DEAL_NAME}}': (meta or {}).get('dealName') or clean_inline(title),
    '{{CLIENT_NAME}}': (meta or {}).get('clientName') or '',
    '{{DATE}}': (meta or {}).get('date') or '',
    '{{TITLE}}': clean_inline(title),
  }
  apply_template_placeholders(doc, replacements)

  marker_paragraph = None
  for paragraph in doc.paragraphs:
    if CONTENT_MARKER in paragraph.text:
      marker_paragraph = paragraph
      break

  body = doc.element.body
  existing_ids = {id(child) for child in body}

  build_content(doc, title, markdown, diagrams, timeline_diagrams, inline_diagram_images, include_cover=False, include_toc=False, meta=meta)

  new_children = [child for child in body if id(child) not in existing_ids]

  if marker_paragraph is not None:
    marker_text = marker_paragraph.text.strip()
    remove_marker_paragraph = marker_text == CONTENT_MARKER
    if not remove_marker_paragraph:
      replacement_text = marker_paragraph.text.replace(CONTENT_MARKER, '').strip()
      clear_paragraph_content(marker_paragraph)
      if replacement_text:
        run = marker_paragraph.add_run(replacement_text)
        set_run_font(run, size=11)
      else:
        remove_marker_paragraph = True

    anchor = marker_paragraph._p
    for child in new_children:
      body.remove(child)
      anchor.addnext(child)
      anchor = child

    if remove_marker_paragraph:
      marker_paragraph._p.getparent().remove(marker_paragraph._p)

  normalize_table_styles(doc)
  return doc


def main():
  payload = json.loads(sys.stdin.read() or '{}')
  markdown = payload.get('markdown') or ''
  output_path = payload['outputPath']
  title = payload.get('title') or 'Proposal'
  template_path = payload.get('templatePath')
  diagrams = payload.get('diagrams') or []
  timeline_diagrams = payload.get('timelineDiagrams') or []
  inline_diagram_images = payload.get('inlineDiagramImages') or []
  meta = {
    'clientName': payload.get('clientName') or '',
    'dealName': payload.get('dealName') or '',
    'date': payload.get('date') or '',
  }

  if template_path and Path(template_path).exists():
    doc = render_into_template(template_path, title, markdown, diagrams, timeline_diagrams, inline_diagram_images, meta=meta)
  else:
    doc = create_content_document(title, markdown, diagrams, timeline_diagrams, inline_diagram_images, meta=meta)

  Path(output_path).parent.mkdir(parents=True, exist_ok=True)
  doc.save(output_path)


if __name__ == '__main__':
  main()
